"""Read the supplied partner layouts without changing the source workbooks."""
from __future__ import annotations

import math
import logging
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import Dataset, Item, Month, Sale, Shipment, Source, SourceRef

MONTHS = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
ROLES = ("monthly_sales", "transactions", "stock_history", "transit", "moq", "seasonality")


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a quantity")
    result = float(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))
    if not math.isfinite(result):
        raise ValueError("nonfinite number")
    return result


def text(value):
    return "" if value is None else str(value).strip()


def month_date(value):
    value = text(value).lower()
    year = re.search(r"20\d{2}", value)
    if not year:
        return None
    for index, prefix in enumerate(MONTHS, 1):
        if value.startswith(prefix):
            return date(int(year.group()), index, 1)
    return None


def dated(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(text(value)[:10], "%d.%m.%Y").date()


def role_for(path):
    name = path.name.lower()
    if "динамика" in name:
        return "transactions"
    if "ежемесячные продажи" in name:
        return "monthly_sales"
    if "ежемесячные остатки" in name:
        return "stock_history"
    if "сезонность" in name:
        return "seasonality"
    if "moq" in name:
        return "moq"
    if "путь" in name or "в пути" in name:
        return "transit"
    return None


def supplier_for(path):
    name = str(path).lower()
    if "system" in name or "syseme" in name:
        return "Systeme Electric"
    if "иэк" in name or "iek" in name:
        return "IEK"
    return None


def load_dataset(paths: list[Path]) -> Dataset:
    files = set()
    issues = []
    for raw in paths:
        path = Path(raw).expanduser()
        if not path.exists():
            issues.append("Источник недоступен: " + path.name)
        elif path.is_dir():
            files.update(p.resolve() for p in path.rglob("*.xlsx") if not p.name.startswith("~$"))
        elif path.suffix.lower() == ".xlsx":
            files.add(path.resolve())
    if not files:
        return Dataset(items=[], issues=issues + ["Не найдено файлов XLSX."])
    items: dict[tuple[str, str], Item] = {}
    sources = []
    seasons = {}
    season_refs = {}
    seen = set()
    failed = set()

    def item_for(supplier, sku, name=""):
        key = supplier, sku
        if key not in items:
            items[key] = Item(supplier=supplier, sku=sku, name=name)
        if not items[key].name:
            items[key].name = name
        return items[key]

    for path in sorted(files, key=lambda p: (ROLES.index(role_for(p)) if role_for(p) else 99, str(p))):
        supplier, role = supplier_for(path), role_for(path)
        if not supplier or not role:
            issues.append("Не распознан тип/поставщик файла: " + path.name)
            continue
        source = Source(file=path.name, supplier=supplier, role=role)
        sources.append(source)
        if (supplier, role) in seen:
            source.status = "invalid_data"
            source.message = "Два источника одной роли: выберите одну версию."
            failed.add(supplier)
            continue
        seen.add((supplier, role))
        wb = None
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb["TDSheet"] if supplier == "Systeme Electric" and role == "transit" else wb.worksheets[0]

            def ref(row, col):
                return SourceRef(file=path.name, sheet=ws.title, cell=f"{get_column_letter(col + 1)}{row}")

            if role == "seasonality":
                values, refs = {}, {}
                for rowno, row in enumerate(ws.iter_rows(min_row=11, max_row=22), 11):
                    label = text(row[1].value).lower()
                    month = next((i for i, prefix in enumerate(MONTHS, 1) if label.startswith(prefix)), None)
                    value = number(row[11].value)
                    if month and value is not None and value > 0:
                        values[month] = value
                        refs[f"season_{month}"] = ref(rowno, 11)
                if len(values) != 12:
                    raise ValueError("seasonal structure")
                seasons[supplier], season_refs[supplier] = values, refs
                source.rows = 12
                continue

            header_row = 2 if role == "transit" and supplier == "Systeme Electric" else 1
            iterator = ws.iter_rows(values_only=True)
            header = []
            for _ in range(header_row):
                header = list(next(iterator))
            labels = [text(v) for v in header]
            if role == "transactions":
                required = ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"]
                if labels[:8] != required:
                    raise ValueError("transaction headers")
                customer_col = next((i for i, label in enumerate(labels) if label in ("client_id", "Обезличенный ID клиента")), None)
                ignored = 0
                for rowno, row in enumerate(iterator, header_row + 1):
                    if not text(row[3]):
                        continue
                    item = item_for(supplier, text(row[3]), text(row[4]))
                    item.unit = text(row[5]) or item.unit
                    if not text(row[2]).startswith("Расходная накладная"):
                        ignored += 1
                        continue
                    try:
                        qty = number(row[7])
                        if qty is None:
                            raise ValueError("missing quantity")
                        item.sales.append(Sale(date=dated(row[0]), quantity=qty, document=text(row[1]),
                                               warehouse=text(row[6]), source=ref(rowno, 7),
                                               customer_id=(text(row[customer_col]) or None) if customer_col is not None else None))
                        source.rows += 1
                    except (ValueError, TypeError):
                        item.invalid = True
                        item.issues.append(f"Некорректная операция: {path.name}, строка {rowno}.")
                source.message = f"Прочитаны накладные; прочих документов: {ignored}. Они не сложены с продажами."
                continue

            if role in ("monthly_sales", "stock_history"):
                if "Номенклатура.Код" not in labels:
                    raise ValueError("missing sku header")
                code_col = labels.index("Номенклатура.Код")
                name_col = labels.index("Номенклатура")
                unit_col = next((i for i, x in enumerate(labels) if x in ("Ед.", "Ед.изм")), None)
                cols = [(i, month_date(v)) for i, v in enumerate(header) if month_date(v)]
                if not cols:
                    raise ValueError("missing month headers")
                local_codes = set()
                for rowno, row in enumerate(iterator, header_row + 1):
                    sku = text(row[code_col])
                    if not sku:
                        continue
                    item = item_for(supplier, sku, text(row[name_col]))
                    if sku in local_codes:
                        item.invalid = True
                        item.issues.append("Повтор кода товара в " + path.name)
                        continue
                    local_codes.add(sku)
                    if unit_col is not None:
                        item.unit = text(row[unit_col]) or item.unit
                    month_map = {m.month: m for m in item.months}
                    for col, month in cols:
                        try:
                            value = number(row[col])
                        except (ValueError, TypeError):
                            item.invalid = True
                            item.issues.append(f"Нечисловое значение: {path.name}, {get_column_letter(col + 1)}{rowno}.")
                            value = None
                        if month not in month_map:
                            month_map[month] = Month(month=month, quantity=None, source=ref(rowno, col))
                        if role == "monthly_sales":
                            month_map[month].quantity = value
                            month_map[month].source = ref(rowno, col)
                        else:
                            month_map[month].opening_stock = value
                            month_map[month].stock_source = ref(rowno, col)
                    item.months = sorted(month_map.values(), key=lambda m: m.month)
                    source.rows += 1
                continue

            code_label = "Код 1с" if "Код 1с" in labels else "Номенклатура.Код"
            if code_label not in labels:
                raise ValueError("missing sku header")
            code_col = labels.index(code_label)
            name_col = next((i for i, x in enumerate(labels) if x in ("Номенклатура", "Наименование")), None)
            article_col = next((i for i, x in enumerate(labels) if x in ("Артикул", "Артикул ИЭК", "Артикул поставщика")), None)
            date_match = re.search(r"\d{2}\.\d{2}\.\d{4}", path.name)
            snapshot = dated(date_match.group()) if date_match else None
            if role == "moq":
                qty_col = next((i for i, x in enumerate(labels) if x in ("Кратность", "Мин. разр. к отгр.")), None)
                if qty_col is None:
                    raise ValueError("missing moq header")
            elif supplier == "Systeme Electric":
                for label in ("Свободный остаток", "Категория 2026", "Кэф. Роста"):
                    if label not in labels:
                        raise ValueError("missing snapshot header")
            local_codes = set()
            for rowno, row in enumerate(iterator, header_row + 1):
                sku = text(row[code_col])
                if not sku:
                    continue
                item = item_for(supplier, sku, text(row[name_col]) if name_col is not None else "")
                if sku in local_codes:
                    item.invalid = True
                    item.issues.append("Повтор кода товара в " + path.name)
                    continue
                local_codes.add(sku)
                if article_col is not None and row[article_col] is not None:
                    item.article = text(row[article_col])
                try:
                    if role == "moq":
                        value = number(row[qty_col])
                        if value is None or value <= 0:
                            item.issues.append("MOQ/кратность отсутствует или не положительна.")
                        else:
                            if supplier == "IEK":
                                item.minimum = value
                            else:
                                item.multiple = value
                            item.sources["moq"] = ref(rowno, qty_col)
                    else:
                        item.transit_known = True
                        item.sources["transit"] = ref(rowno, code_col)
                        if "БУХТАМИ" in item.name.upper():
                            item.unit_conversion_required = True
                        shipment_columns = 0
                        missing_shipment_quantity = False
                        for col, label in enumerate(labels):
                            if "поступление до" in label:
                                arrival_match = re.search(r"поступление до (\d{2}\.\d{2}\.\d{4})", label)
                            elif "в пути" in label.lower():
                                arrival_match = re.search(r"(\d{2}\.\d{2})", label)
                            else:
                                continue
                            shipment_columns += 1
                            qty = number(row[col])
                            if qty is None:
                                missing_shipment_quantity = True
                                continue
                            if qty == 0:
                                continue
                            arrival_text = arrival_match.group(1) if arrival_match else None
                            if arrival_text and len(arrival_text) == 5 and snapshot:
                                arrival_text += "." + str(snapshot.year)
                            arrival = dated(arrival_text) if arrival_text else None
                            item.shipments.append(Shipment(quantity=qty, arrival=arrival, source=ref(rowno, col)))
                        if not shipment_columns:
                            raise ValueError("missing shipment headers")
                        if missing_shipment_quantity:
                            item.transit_known = False
                            item.issues.append("В таблице поставок есть пустые количества; они не приравнены к нулю.")
                        if supplier == "Systeme Electric":
                            col = labels.index("Свободный остаток")
                            stock = number(row[col])
                            if stock is not None and stock < 0:
                                raise ValueError("negative free stock")
                            item.stock, item.stock_date = stock, snapshot
                            item.sources["stock"] = ref(rowno, col)
                            c = labels.index("Категория 2026")
                            item.category = text(row[c]) or None
                            item.sources["category"] = ref(rowno, c)
                            c = labels.index("Кэф. Роста")
                            rate = number(row[c])
                            if rate is not None and rate < -1:
                                raise ValueError("invalid growth")
                            item.growth_rate = rate
                            item.sources["growth"] = ref(rowno, c)
                    source.rows += 1
                except (ValueError, TypeError):
                    item.invalid = True
                    item.issues.append(f"Некорректное значение: {path.name}, строка {rowno}.")
        except OSError:
            source.status, source.message = "source_unavailable", "Нет доступа к файлу."
            failed.add(supplier)
        except Exception as exc:
            logging.getLogger(__name__).error("XLSX import failed (%s, %s): %s", supplier, role, type(exc).__name__)
            source.status, source.message = "invalid_data", "Не удалось прочитать структуру или значения XLSX."
            failed.add(supplier)
        finally:
            if wb:
                wb.close()

    for item in items.values():
        item.seasonality = seasons.get(item.supplier, {}).copy()
        item.sources.update(season_refs.get(item.supplier, {}))
        if item.supplier in failed:
            item.invalid = True
            item.issues.append("Один из источников поставщика некорректен или недоступен.")
        missing = [r for r in ROLES if (item.supplier, r) not in seen]
        if missing:
            item.issues.append("Не переданы источники: " + ", ".join(missing))
    for source in sources:
        if source.status != "ok":
            issues.append(source.file + ": " + source.message)
    return Dataset(items=sorted(items.values(), key=lambda i: i.key), sources=sources, issues=issues)
